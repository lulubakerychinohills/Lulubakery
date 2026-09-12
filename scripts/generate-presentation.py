#!/usr/bin/env python3
"""Elegant final presentation for Lulu Bakery — refined type & backgrounds."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt, Emu

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs" / "ppt-assets"
OUT = ROOT / "docs" / "Lulu-Bakery-Final-Presentation.pptx"

# Palette — quiet, editorial
INK = RGBColor(0x2A, 0x24, 0x20)
MUTED = RGBColor(0x7A, 0x6F, 0x66)
BRAND = RGBColor(0x4A, 0x3C, 0x35)
ACCENT = RGBColor(0xA8, 0x8F, 0x78)
CREAM = RGBColor(0xFB, 0xF8, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD4, 0xCB, 0xC0)
IVORY = RGBColor(0xF7, 0xF3, 0xEE)

TITLE_FONT = "Georgia"
BODY_FONT = "Calibri"
LIGHT_FONT = "Calibri Light"


def run_style(run, *, size=18, bold=False, color=INK, font=BODY_FONT, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font


def add_bg(slide, name: str):
    slide.shapes.add_picture(
        str(ASSETS / name),
        Inches(0),
        Inches(0),
        width=Inches(13.333),
        height=Inches(7.5),
    )


def add_rect(slide, left, top, width, height, fill, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    return shape


def add_line(slide, left, top, width):
    """Thin horizontal accent rule."""
    return add_rect(slide, left, top, width, 0.015, ACCENT)


def add_vbar(slide, left=0.55, top=0.9, height=5.6):
    return add_rect(slide, left, top, 0.03, height, ACCENT)


def add_textbox(slide, left, top, width, height, paragraphs, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(paragraphs):
        text, size, bold, color, space_after = item[:5]
        font = item[5] if len(item) > 5 else BODY_FONT
        italic = item[6] if len(item) > 6 else False
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = text
        run_style(r, size=size, bold=bold, color=color, font=font, italic=italic)
    return box


def add_bullets(slide, left, top, width, height, lines):
    paras = []
    for line in lines:
        paras.append((line, 16, False, INK, 14, BODY_FONT, False))
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        p.space_before = Pt(2)
        # accent dash instead of heavy bullet
        r0 = p.add_run()
        r0.text = "–  "
        run_style(r0, size=15, color=ACCENT, font=TITLE_FONT)
        r1 = p.add_run()
        r1.text = line
        run_style(r1, size=16, color=INK, font=BODY_FONT)
    return box


def add_footer(slide, page: int, total: int = 12, light=False):
    color = RGBColor(0xB0, 0xA6, 0x9C) if light else MUTED
    add_textbox(
        slide,
        0.7,
        7.05,
        10.5,
        0.28,
        [(f"Lulu Bakery    {page} / {total}", 10, False, color, 0, LIGHT_FONT, False)],
    )


def add_picture_framed(slide, name, left, top, width):
    # soft frame
    path = ASSETS / name
    # slightly larger backing
    add_rect(slide, left - 0.06, top - 0.06, width + 0.12, width + 0.12, LINE)
    pic = slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))
    return pic


def content_slide(prs, page, bg, kicker, title, bullets, image=None):
    blank = prs.slide_layouts[6]
    s = prs.slides.add_slide(blank)
    add_bg(s, bg)

    # soft content veil (not a chunky card)
    add_rect(s, 0.35, 0.3, 12.6, 6.7, IVORY)

    add_vbar(s, 0.7, 0.85, 5.5)

    add_textbox(
        s,
        1.0,
        0.7,
        10,
        0.35,
        [(kicker.upper(), 11, False, ACCENT, 0, BODY_FONT, False)],
    )
    add_textbox(
        s,
        1.0,
        1.05,
        10.5,
        0.6,
        [(title, 28, False, BRAND, 0, TITLE_FONT, False)],
    )
    add_line(s, 1.0, 1.75, 1.4)

    text_width = 7.0 if image else 11.0
    add_bullets(s, 1.0, 2.1, text_width, 4.4, bullets)

    if image:
        add_picture_framed(s, image, 8.7, 2.15, 3.6)

    add_footer(s, page)
    return s


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # ---- 1 Title ----
    s = prs.slides.add_slide(blank)
    add_bg(s, "ppt-bg-espresso.jpg")
    add_textbox(
        s,
        1.1,
        2.0,
        8.5,
        0.35,
        [("WEB DEVELOPMENT  ·  FINAL PROJECT", 11, False, ACCENT, 0, BODY_FONT, False)],
    )
    add_textbox(
        s,
        1.1,
        2.45,
        9.5,
        1.1,
        [("Lulu Bakery", 48, False, CREAM, 6, TITLE_FONT, False)],
    )
    add_line(s, 1.1, 3.55, 1.6)
    add_textbox(
        s,
        1.1,
        3.8,
        8.5,
        1.4,
        [
            ("A custom cake ordering website", 18, False, RGBColor(0xD8, 0xCF, 0xC5), 8, BODY_FONT, False),
            ("for a local bakery in Chino Hills", 18, False, RGBColor(0xD8, 0xCF, 0xC5), 14, BODY_FONT, False),
            ("Next.js  ·  React  ·  lulubakerychinohills.com", 13, False, RGBColor(0xA8, 0x9C, 0x92), 0, LIGHT_FONT, False),
        ],
    )
    # small logo, quiet
    add_rect(s, 10.35, 2.35, 1.85, 1.85, RGBColor(0x3A, 0x30, 0x2A))
    slide_pic = s.shapes.add_picture(
        str(ASSETS / "logo-avatar.jpg"),
        Inches(10.45),
        Inches(2.45),
        width=Inches(1.65),
    )
    add_footer(s, 1, light=True)

    # ---- content slides ----
    content_slide(
        prs,
        2,
        "ppt-bg-ivory.jpg",
        "Overview",
        "What this project is",
        [
            "This site is for Lulu Bakery in Chino Hills.",
            "They make custom cakes. People need to see styles and place an order online.",
            "Most visitors are probably parents planning a birthday or family event.",
            "They browse on the phone first, then finish the form later.",
            "So the flow stays simple: look, choose, submit.",
        ],
        image="hero-cupcake.jpg",
    )

    content_slide(
        prs,
        3,
        "ppt-bg-stone.jpg",
        "Architecture",
        "Site structure",
        [
            "Home — cake gallery and order form",
            "Category pages — men / women / kids / other",
            "Sweet — dessert menu, then a photos page",
            "About — who they are and pickup info",
            "Privacy — required before someone submits an order",
            "The main nav stays short on purpose.",
        ],
    )

    content_slide(
        prs,
        4,
        "ppt-bg-ivory.jpg",
        "Design",
        "Look and feel",
        [
            "Warm brown on cream — quiet, not loud.",
            "I checked contrast so body text stays readable.",
            "Same buttons and cards across pages.",
            "Geist Sans through next/font on the site.",
            "Spacing comes from a small set of CSS variables.",
        ],
        image="logo-avatar.jpg",
    )

    content_slide(
        prs,
        5,
        "ppt-bg-stone.jpg",
        "CSS",
        "How the layout is built",
        [
            "box-sizing: border-box everywhere.",
            "Header and filter chips use Flexbox.",
            "The cake grid and order form use CSS Grid.",
            "Nav is sticky. The success popup is fixed.",
            "There’s a basic print stylesheet too.",
        ],
    )

    content_slide(
        prs,
        6,
        "ppt-bg-ivory.jpg",
        "Responsive",
        "How it adapts",
        [
            "Mobile first, then sm / md / lg breakpoints.",
            "Phone: one column. Desktop: three cake cards.",
            "The form goes to two columns on wider screens.",
            "Images use lazy loading and sizes.",
            "I’ll resize the browser in the demo.",
        ],
        image="hero-cupcake.jpg",
    )

    content_slide(
        prs,
        7,
        "ppt-bg-stone.jpg",
        "JavaScript",
        "Interactivity",
        [
            "Gallery search waits 300ms before filtering.",
            "Category buttons also update the URL.",
            "The form checks email, pickup time, and privacy.",
            "After submit, a dialog appears. Esc closes it.",
            "Native HTML required still helps if JS is slow.",
        ],
    )

    content_slide(
        prs,
        8,
        "ppt-bg-ivory.jpg",
        "Accessibility",
        "What I paid attention to",
        [
            "Skip link to jump past the header.",
            "Real <nav>, plus breadcrumbs on deeper pages.",
            "Focus outline shows when you tab.",
            "Dialog uses role=\"dialog\" and labels.",
            "Tests include axe; I’ll also show Lighthouse / WAVE.",
        ],
    )

    content_slide(
        prs,
        9,
        "ppt-bg-stone.jpg",
        "Build",
        "Stack and performance",
        [
            "Next.js, React, TypeScript, Tailwind.",
            "Products and images live in Supabase.",
            "Orders go out by SMTP email.",
            "WebP where I can; uploads compressed with sharp.",
            "Deployed on Vercel.",
        ],
        image="logo-avatar.jpg",
    )

    content_slide(
        prs,
        10,
        "ppt-bg-ivory.jpg",
        "Process",
        "Testing and docs",
        [
            "npm test — 13 tests.",
            "Search, form checks, skip link, breadcrumbs, axe smoke.",
            "design-dossier.md for the design write-up.",
            "accessibility-report.md for the checklist.",
            "README has setup and where the docs are.",
        ],
    )

    content_slide(
        prs,
        11,
        "ppt-bg-stone.jpg",
        "Reflection",
        "What was hard / what I learned",
        [
            "The shop already worked — the class still needed clearer a11y and tests.",
            "The success popup looked fine, but needed proper dialog markup.",
            "Pulling search into a plain function made testing easier.",
            "One shared header beat copying the same nav three times.",
            "Landmarks are easier if you set them up early.",
        ],
    )

    # ---- 12 Thanks ----
    s = prs.slides.add_slide(blank)
    add_bg(s, "ppt-bg-espresso.jpg")
    add_textbox(
        s,
        1.1,
        2.4,
        11,
        0.9,
        [("Thank you", 48, False, CREAM, 0, TITLE_FONT, False)],
    )
    add_line(s, 1.1, 3.4, 1.4)
    add_textbox(
        s,
        1.1,
        3.7,
        10,
        1.6,
        [
            ("Live demo next — layout, search, form, keyboard nav", 16, False, RGBColor(0xD0, 0xC6, 0xBC), 12, BODY_FONT, False),
            ("lulubakerychinohills.com", 15, False, ACCENT, 14, BODY_FONT, False),
            ("Questions are welcome", 14, False, RGBColor(0xA8, 0x9C, 0x92), 0, LIGHT_FONT, True),
        ],
    )
    add_footer(s, 12, light=True)

    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
